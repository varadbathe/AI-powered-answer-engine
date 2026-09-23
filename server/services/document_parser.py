import io
import re
from typing import Any
import docx
from pypdf import PdfReader
from pypdf.errors import PdfReadError


class DocumentParserError(Exception):
    pass


class EmptyDocumentError(DocumentParserError):
    pass


class CorruptedDocumentError(DocumentParserError):
    pass


class DocumentParser:
    @staticmethod
    def clean_text(text: str) -> str:
        """
        Cleans text: removes control characters (preserving newlines and tabs),
        normalizes excessive whitespace and blank lines.
        """
        if not text:
            return ""

        # Remove null bytes and unprintable control characters
        text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", "", text)

        # Normalize carriage returns
        text = text.replace("\r\n", "\n").replace("\r", "\n")

        # Replace excessive consecutive newlines with at most two
        text = re.sub(r"\n{3,}", "\n\n", text)

        # Replace tabs and multiple spaces with single space
        text = re.sub(r"[ \t]+", " ", text)

        return text.strip()

    @classmethod
    def parse_pdf(cls, file_bytes: bytes) -> list[dict[str, Any]]:
        """
        Parses PDF using pypdf, extracting text per page and tracking page numbers.
        """
        try:
            reader = PdfReader(io.BytesIO(file_bytes))
        except PdfReadError as e:
            raise CorruptedDocumentError(f"Corrupted or invalid PDF: {e}")
        except Exception as e:
            raise CorruptedDocumentError(f"Failed to read PDF: {e}")

        if len(reader.pages) == 0:
            raise EmptyDocumentError("PDF file has 0 pages")

        segments = []
        total_text_len = 0

        for page_idx, page in enumerate(reader.pages):
            try:
                page_text = page.extract_text() or ""
            except Exception as e:
                page_text = ""

            cleaned = cls.clean_text(page_text)
            if cleaned:
                total_text_len += len(cleaned)
                segments.append({
                    "text": cleaned,
                    "page_number": page_idx + 1,
                    "section_title": None,
                })

        if total_text_len == 0:
            raise EmptyDocumentError("PDF contains no extractable text")

        return segments

    @classmethod
    def parse_docx(cls, file_bytes: bytes) -> list[dict[str, Any]]:
        """
        Parses DOCX using python-docx, detecting headings as section_title
        and extracting paragraphs and tables.
        """
        try:
            doc = docx.Document(io.BytesIO(file_bytes))
        except Exception as e:
            raise CorruptedDocumentError(f"Corrupted or invalid DOCX document: {e}")

        segments = []
        current_section = None
        current_paras: list[str] = []
        total_text_len = 0

        def flush_section():
            nonlocal current_paras, total_text_len
            combined = "\n\n".join(current_paras).strip()
            cleaned = cls.clean_text(combined)
            if cleaned:
                total_text_len += len(cleaned)
                segments.append({
                    "text": cleaned,
                    "page_number": None,
                    "section_title": current_section,
                })
            current_paras = []

        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue

            style_name = getattr(p.style, "name", "")
            if style_name and "heading" in style_name.lower():
                flush_section()
                current_section = text
            else:
                current_paras.append(text)

        # Extract table text
        for table in doc.tables:
            table_rows = []
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    table_rows.append(" | ".join(cells))
            if table_rows:
                current_paras.append("\n".join(table_rows))

        flush_section()

        if total_text_len == 0:
            raise EmptyDocumentError("DOCX contains no extractable text")

        return segments

    @classmethod
    def parse_txt_or_markdown(cls, file_bytes: bytes, is_markdown: bool = False) -> list[dict[str, Any]]:
        """
        Parses TXT or Markdown file, detecting Markdown headers as section_title.
        """
        try:
            raw_text = file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            try:
                raw_text = file_bytes.decode("latin-1")
            except Exception as e:
                raise CorruptedDocumentError(f"Failed to decode text file: {e}")

        cleaned_full = cls.clean_text(raw_text)
        if not cleaned_full:
            raise EmptyDocumentError("Text document is empty")

        if not is_markdown:
            return [{
                "text": cleaned_full,
                "page_number": None,
                "section_title": None,
            }]

        # For Markdown, split by headers (# Header)
        segments = []
        lines = raw_text.splitlines()
        current_section = None
        current_lines: list[str] = []
        total_text_len = 0

        def flush_md_section():
            nonlocal current_lines, total_text_len
            combined = "\n".join(current_lines).strip()
            cleaned = cls.clean_text(combined)
            if cleaned:
                total_text_len += len(cleaned)
                segments.append({
                    "text": cleaned,
                    "page_number": None,
                    "section_title": current_section,
                })
            current_lines = []

        header_regex = re.compile(r"^(#{1,6})\s+(.+)$")
        for line in lines:
            match = header_regex.match(line.strip())
            if match:
                flush_md_section()
                current_section = match.group(2).strip()
            else:
                current_lines.append(line)

        flush_md_section()

        if total_text_len == 0:
            # If no sections with headers or all empty, return the cleaned full text
            return [{
                "text": cleaned_full,
                "page_number": None,
                "section_title": None,
            }]

        return segments

    @classmethod
    def parse(cls, file_bytes: bytes, file_type: str) -> list[dict[str, Any]]:
        """
        Dispatches parsing based on file_type ('pdf', 'docx', 'txt', 'md').
        Returns list of segment dictionaries with keys: 'text', 'page_number', 'section_title'.
        """
        ft = file_type.lower().lstrip(".")
        if ft == "pdf":
            return cls.parse_pdf(file_bytes)
        elif ft == "docx":
            return cls.parse_docx(file_bytes)
        elif ft == "md":
            return cls.parse_txt_or_markdown(file_bytes, is_markdown=True)
        elif ft == "txt":
            return cls.parse_txt_or_markdown(file_bytes, is_markdown=False)
        else:
            raise DocumentParserError(f"Unsupported file type for parsing: {file_type}")
